#!/usr/bin/env python3
"""Export the Editorial Studio database to the Hnahsin content Google Sheet.

Writes an .xlsx with the tabs content_studio/Pack.js expects (Thumal, Zawhna,
Sentence, Game thu), the uploaded pictures as files for the “Hnahsin
thlalak” Drive folder, and the same rows as JSON for checking the export with
Pack.js:

    python3 scripts/export_studio_to_sheet.py OUT_DIR [DATABASE]

Needs openpyxl. Each item exports the text the app shows today: its
published revision, or its latest revision if it was never published.
Published items whose meaning, English or example is missing (the app drops
them now) come out as “Endik mek” with a note, so the first Publish from the
sheet is not blocked. Notes from content/review sheets go to Hriattirna.
"""

from __future__ import annotations

import base64
import csv
import json
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
GAMES = [
    ("picture_match", "Picture Match"), ("spelling", "Spelling"), ("word_search", "Word Search"),
    ("word_chain", "Word Chain"), ("tawng_upa", "Tawng Upa"), ("crossword", "Crossword"),
    ("thumal_kawp", "Thumal Kawp"), ("sentence_builder", "Sentence Builder"),
]
CATEGORY_LABELS = {"chhungkua": "Chhungkua", "sikul": "Sikul", "nungcha": "Nungcha",
                   "khawvel": "Khawvel", "nunphung": "Nunphung", "thiltih": "Thiltih"}
LIVE, REVIEW, DROPPED = "Chhuah", "Endik mek", "Paih"
EXTENSIONS = {"image/png": "png", "image/jpeg": "jpg", "image/webp": "webp"}

WORD_HEADINGS = ["ID", "Thumal", "Awmzia (Mizo)", "English", "Entirna", "Emoji", "Thlalak", "Pawl", "Level",
                 *[title for _, title in GAMES], "Dinhmun", "Hriattirna"]
QUESTION_HEADINGS = ["ID", "Zawhna", "Chhanna 1", "Chhanna 2", "Chhanna 3", "Chhanna 4", "Chhanna dik (1-4)",
                     "Hrilhfiahna", "Emoji", "Level", "Dinhmun", "Hriattirna"]
SENTENCE_HEADINGS = ["ID", "Sentence (Mizo)", "English", "Level", "Dinhmun", "Hriattirna"]
GAME_TEXT_HEADINGS = ["ID", "Game", "Hming", "Hming hnuai thu", "Zawhna thu", "Hint", "Khelh dan 1", "Khelh dan 2",
                      "Khelh dan 3", "A dik a nih chuan", "A dik loh chuan", "Dinhmun", "Hriattirna"]


def seed_categories() -> dict[str, str]:
    """The app's seed_category → category table, read from its source."""
    source = (ROOT / "lib/features/content_sync/domain/delivery_models.dart").read_text()
    block = source.split("_seedCategoryFallback = <String, String>{", 1)[1].split("};", 1)[0]
    return dict(re.findall(r"'([^']+)':\s*'([^']+)'", block))


def query(database: str, sql: str) -> list:
    out = subprocess.run(["psql", "-d", database, "-At", "-c", sql], check=True, capture_output=True, text=True).stdout
    return json.loads(out) if out.strip() else []


def items(database: str) -> list:
    """Every item with the body the app shows: published revision first."""
    return query(database, """
        select coalesce(json_agg(json_build_object('id', i.stable_id, 'status', i.status, 'body', r.body,
               'pending', i.published_revision_id is not null and i.status <> 3) order by i.stable_id), '[]')
        from content_items i
        join lateral (
          select body from content_revisions r
          where r.id = i.published_revision_id
             or (i.published_revision_id is null and r.content_item_id = i.id)
          order by r.number desc limit 1) r on true""")


def status_of(item: dict) -> str:
    return {3: LIVE, 4: DROPPED}.get(item["status"], LIVE if item["pending"] else REVIEW)


def review_notes() -> dict[str, str]:
    notes: dict[str, list[str]] = {}
    check = ROOT / "content/review/language_check_2026-10-01.csv"
    if check.exists():
        for row in csv.DictReader(check.open()):
            suggestion = row["draft_english_meaning"] if row["draft_source"].startswith("English suggested") else ""
            notes.setdefault(row["stable_id"], []).append(
                f"Endik tur ({row['confidence']}): {row['issue']}" + (f" English rawt: {suggestion}" if suggestion else ""))
    emoji = ROOT / "content/review/emoji_suggestions_2026-10-01.json"
    if emoji.exists():
        for stable_id, value in json.loads(emoji.read_text())["corrections"].items():
            notes.setdefault(stable_id, []).append(f"Emoji rawt: {value}")
    corrections = ROOT / "content/review/emoji_corrections_2026-10-01.json"
    if corrections.exists():
        for stable_id, value in json.loads(corrections.read_text())["corrections"].items():
            notes.setdefault(stable_id, []).append(
                f"Emoji a dik lo: {value} hmang rawh" if value else "Emoji a dik lo: paih rawh")
    return {key: " | ".join(value) for key, value in notes.items()}


def drafts() -> dict[str, dict]:
    """Meaning drafts for words that have none yet (filled in, still Endik mek)."""
    out = {}
    for name in ["missing_meanings_2026-09-28.csv", "language_check_2026-10-01.csv"]:
        path = ROOT / "content/review" / name
        if not path.exists():
            continue
        for row in csv.DictReader(path.open()):
            if name.startswith("language_check") and row["what_to_check"] != "all":
                continue
            out[row["stable_id"]] = row
    return out


def export(out_dir: Path, database: str) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    pictures = out_dir / "Hnahsin thlalak"
    pictures.mkdir(exist_ok=True)
    seeds = seed_categories()
    notes = review_notes()
    word_drafts = drafts()
    fixed_questions = json.loads((ROOT / "content/tawng_upa/default_questions.json").read_text())["questions"]
    images = {row["checksum"]: row for row in query(database, """
        select coalesce(json_agg(json_build_object('checksum', checksum_sha256, 'type', content_type,
               'data', encode(data, 'base64'))), '[]') from word_images""")}

    tabs: dict[str, list[dict]] = {"words": [], "questions": [], "sentences": [], "gameText": []}
    for item in items(database):
        body, sid, status = item["body"], item["id"], status_of(item)
        note = [notes[sid]] if sid in notes else []
        if item["pending"]:
            note.append("Studio-ah a thlak tur endik mek a awm (Studio-a version thar hi a lut lo).")
        kind = body.get("type")
        if kind == "word":
            content, learning = body.get("content", {}), body.get("learning", {})
            category = next((seeds[c] for c in learning.get("categories", []) if c in seeds), "")
            row = {
                "id": sid, "word": content.get("canonical_form") or "", "meaning": content.get("definition_mizo") or "",
                "english": (content.get("glosses") or {}).get("en") or "", "example": content.get("example_mizo") or "",
                "emoji": content.get("emoji") or "", "picture": "", "category": CATEGORY_LABELS.get(category, ""),
                "level": learning.get("difficulty") or "", "status": status,
            }
            for game, _ in GAMES:
                row[game] = game in (learning.get("game_modes") or [])
            image = content.get("image")
            if image and image.get("checksum") in images:
                stored = images[image["checksum"]]
                name = f"{sid}.{EXTENSIONS[stored['type']]}"
                (pictures / name).write_bytes(base64.b64decode(stored["data"]))
                row["picture"] = name
            missing = [label for key, label in [("meaning", "Awmzia"), ("english", "English"), ("example", "Entirna")]
                       if not row[key].strip()]
            draft = word_drafts.get(sid)
            if draft and missing:
                row["meaning"] = row["meaning"] or draft.get("draft_mizo_meaning", "")
                row["english"] = row["english"] or draft.get("draft_english_meaning", "")
                row["example"] = row["example"] or draft.get("draft_example", "")
                row["emoji"] = row["emoji"] or draft.get("draft_emoji", "")
                note.append("Awmzia/English/Entirna hi Claude draft a ni: endik la, a dik chuan Dinhmun “Chhuah” ah thlak rawh.")
            invalid = missing or not row["category"] or not (isinstance(row["level"], int) and 1 <= row["level"] <= 7)
            if status == LIVE and invalid:
                row["status"] = REVIEW
                note.append("App-ah a lang lo: " + ", ".join(missing or ["Pawl/Level"]) + " a ruak.")
            row["note"] = " | ".join(note)
            tabs["words"].append(row)
        elif kind == "question":
            options = body.get("options", [])
            row = {"id": sid, "prompt": body.get("prompt_mizo", ""), "answer": options.index(body["answer"]) + 1,
                   "explanation": body.get("explanation_mizo", ""), "emoji": body.get("emoji", ""),
                   "level": body.get("difficulty", 1), "status": status}
            for n, option in enumerate(options[:4], 1):
                row[f"option{n}"] = option
            # The importer named these question.tawng-upa.001… in the order of
            # content/tawng_upa/default_questions.json, which has the fixed text.
            number = re.fullmatch(r"question\.tawng-upa\.(\d{3})", sid)
            fixed = fixed_questions[int(number.group(1)) - 1] if number and int(number.group(1)) <= len(fixed_questions) else None
            if fixed and (fixed["prompt_mizo"], fixed["options"]) != (body.get("prompt_mizo"), options):
                row.update({"prompt": fixed["prompt_mizo"], "answer": fixed["options"].index(fixed["answer"]) + 1,
                            "explanation": fixed["explanation_mizo"]})
                for n, option in enumerate(fixed["options"], 1):
                    row[f"option{n}"] = option
                note.append("Claude-in a siam ṭha (chhanna dik hian thumal a sawi lo tawh): endik rawh.")
            row["note"] = " | ".join(note)
            tabs["questions"].append(row)
        elif kind == "sentence":
            content, learning = body.get("content", body), body.get("learning", {})
            tq = re.search(r"\d+", str(learning.get("tq_level", "")))
            tabs["sentences"].append({
                "id": sid, "text": content.get("text_mizo", ""), "english": content.get("english_support", ""),
                "level": learning.get("difficulty") or (int(tq.group()) + 1 if tq else 1), "status": status,
                "note": " | ".join(note)})
        elif kind == "game_copy":
            steps = (body.get("instructions") or []) + ["", "", ""]
            tabs["gameText"].append({
                "id": sid, "game": body.get("game_id", ""), "title": body.get("title", ""),
                "subtitle": body.get("subtitle", ""), "prompt": body.get("prompt", ""), "hint": body.get("hint", ""),
                "step1": steps[0], "step2": steps[1], "step3": steps[2],
                "correct_feedback": body.get("correct_feedback", ""), "retry_feedback": body.get("retry_feedback", ""),
                # Game text only changes wording, so the reviewed defaults go live as they are.
                "status": LIVE if status != DROPPED else DROPPED, "note": " | ".join(note)})

    write_xlsx(out_dir / "Hnahsin content.xlsx", tabs)
    (out_dir / "rows.json").write_text(json.dumps(tabs, ensure_ascii=False, indent=1))
    for key, rows in tabs.items():
        live = sum(row["status"] == LIVE for row in rows)
        print(f"{key}: {len(rows)} rows, {live} Chhuah")
    print(f"pictures: {len(list(pictures.iterdir()))}")


def write_xlsx(path: Path, tabs: dict[str, list[dict]]) -> None:
    book = Workbook()
    book.remove(book.active)
    header = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E5F")
    keys = {
        "Thumal": ("words", WORD_HEADINGS, ["id", "word", "meaning", "english", "example", "emoji", "picture", "category",
                                            "level", *[g for g, _ in GAMES], "status", "note"]),
        "Zawhna": ("questions", QUESTION_HEADINGS, ["id", "prompt", "option1", "option2", "option3", "option4", "answer",
                                                    "explanation", "emoji", "level", "status", "note"]),
        "Sentence": ("sentences", SENTENCE_HEADINGS, ["id", "text", "english", "level", "status", "note"]),
        "Game thu": ("gameText", GAME_TEXT_HEADINGS, ["id", "game", "title", "subtitle", "prompt", "hint", "step1", "step2",
                                                      "step3", "correct_feedback", "retry_feedback", "status", "note"]),
    }
    for title, (tab, headings, columns) in keys.items():
        sheet = book.create_sheet(title)
        sheet.append(headings)
        for cell in sheet[1]:
            cell.font, cell.fill = header, fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in tabs[tab]:
            sheet.append([row.get(key, "") for key in columns])
        sheet.freeze_panes = "C2"
        last = max(sheet.max_row, 2000)
        letter = {h: sheet.cell(1, i + 1).column_letter for i, h in enumerate(headings)}
        status = DataValidation(type="list", formula1=f'"{LIVE},{REVIEW},{DROPPED}"', allow_blank=True)
        status.add(f"{letter['Dinhmun']}2:{letter['Dinhmun']}{last}")
        sheet.add_data_validation(status)
        if "Level" in letter:
            level = DataValidation(type="list", formula1='"1,2,3,4,5,6,7"', allow_blank=True)
            level.add(f"{letter['Level']}2:{letter['Level']}{last}")
            sheet.add_data_validation(level)
        for heading in headings:
            sheet.column_dimensions[letter[heading]].width = {
                "ID": 22, "Awmzia (Mizo)": 45, "Entirna": 40, "Hriattirna": 60, "English": 28, "Hrilhfiahna": 45,
                "Sentence (Mizo)": 40, "Zawhna": 36}.get(heading, 14)
    book.save(path)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    export(Path(sys.argv[1]), sys.argv[2] if len(sys.argv) > 2 else "thumal_quest_editorial_development")
